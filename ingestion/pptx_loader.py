# ingestion/pptx_loader.py
import logging
from pptx import Presentation

logger = logging.getLogger(__name__)


def load_pptx(file_path):
    """Load text from a PowerPoint (.pptx) file, slide by slide."""
    try:
        prs = Presentation(file_path)
    except Exception as e:
        logger.error(f"Failed to read PPTX '{file_path}': {e}")
        return []

    slides_text = []

    for slide_num, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        texts.append(line)

        if texts:
            slides_text.append({
                "content": "\n".join(texts),
                "metadata": {
                    "source": file_path,
                    "type": "pptx",
                    "slide": slide_num
                }
            })

    if not slides_text:
        logger.warning(f"No text content found in '{file_path}'")

    return slides_text
